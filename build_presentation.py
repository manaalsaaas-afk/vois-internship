# -*- coding: utf-8 -*-
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BULLET = "\u2022 "

def build_presentation():
    input_path = 'VOIS_Major_Project_PPT_Submission_Template (2).pptx'
    output_path = 'VOIS_Major_Project_Seasonal_Agriculture_Performance_Analysis.pptx'
    
    prs = pptx.Presentation(input_path)
    print(f"Loaded template with {len(prs.slides)} slides.")
    
    # Colors
    COLOR_RED = RGBColor(230, 0, 0)       # VOIS Signature Red
    COLOR_DARK = RGBColor(30, 41, 59)     # Dark Slate / Charcoal
    COLOR_MUTED = RGBColor(100, 116, 139) # Slate Gray
    COLOR_CARD_BG = RGBColor(248, 249, 250)
    COLOR_BORDER = RGBColor(200, 205, 210)
    
    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Seasonal Agriculture Performance Analysis"
            p0.font.size = Pt(32)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_RED
            p1 = tf.add_paragraph()
            p1.text = "VOIS AICTE Internship Major Project | Batch 1 (2026-2027)"
            p1.font.size = Pt(16)
            p1.font.color.rgb = COLOR_MUTED
            shape.left = Inches(0.8)
            shape.top = Inches(2.6)
            shape.width = Inches(10.5)
            shape.height = Inches(1.3)
            
        elif shape.name == 'Text Placeholder 1' and shape.left < 0: # The misaligned template box
            shape.left = Inches(0.8)
            shape.top = Inches(4.2)
            shape.width = Inches(5.5)
            shape.height = Inches(1.1)
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Student Name: Manal Sas"
            p0.font.size = Pt(18)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_DARK
            p1 = tf.add_paragraph()
            p1.text = "College Name: [Your College / University Name]"
            p1.font.size = Pt(15)
            p1.font.color.rgb = COLOR_DARK
            
        elif shape.name == 'TextBox 2' and shape.has_text_frame:
            shape.left = Inches(0.8)
            shape.top = Inches(5.4)
            shape.width = Inches(5.5)
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "AICTE STU ID: [Enter AICTE Student ID]"
            p0.font.size = Pt(15)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_RED
            
    print("Slide 1 formatted.")

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.text = "PROBLEM STATEMENT"
            shape.text_frame.paragraphs[0].font.size = Pt(30)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
        elif shape.name == 'Text Placeholder 1' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            shape.width = Inches(7.5)
            shape.top = Inches(1.6)
            shape.height = Inches(4.8)
            
            points = [
                ("Agricultural Volatility across Seasons: ", "Farming activities and crop yields are heavily governed by seasonal shifts in temperature, monsoon rainfall, atmospheric humidity, and soil moisture across Kharif, Rabi, and Zaid seasons."),
                ("The Information Gap: ", "Raw agricultural records are disconnected and multi-dimensional. Without rigorous analytics, stakeholders cannot discern how seasonal dynamics alter yield efficiency, resource depletion, or crop loss."),
                ("Analytical Research Problem: ", "Systematically evaluate 4,000 farm profiles across 8 Indian states to uncover empirical seasonal patterns, environmental risk thresholds, irrigation productivity, and financial outcomes."),
                ("Strategic & Economic Importance: ", "Deliver evidence-based intelligence to support optimized seasonal crop planning, water conservation, disease mitigation, and agricultural policy planning.")
            ]
            for idx, (title, desc) in enumerate(points):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = BULLET + title
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = COLOR_RED
                p.space_after = Pt(2)
                
                p_desc = tf.add_paragraph()
                p_desc.text = "   " + desc
                p_desc.font.size = Pt(12)
                p_desc.font.color.rgb = COLOR_DARK
                p_desc.space_after = Pt(10)
    print("Slide 2 formatted.")

    # -------------------------------------------------------------
    # SLIDE 3: Project Description
    # -------------------------------------------------------------
    slide3 = prs.slides[2]
    for shape in slide3.shapes:
        if shape.name == 'Title 2' and shape.has_text_frame:
            shape.text_frame.clear()
            p0 = shape.text_frame.paragraphs[0]
            p0.text = "Project Description & Methodology"
            p0.font.size = Pt(28)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.6)
            shape.width = Inches(10)
            
    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.9))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    
    sections = [
        ("1. Dataset Scope & Architecture:", 
         "Analyzed 4,000 farm entries spanning 8 Indian States (AP, MH, TS, KA, GJ, TN, PB, MP) and 10 districts across 8 crops and 3 seasons. The dataset integrates 28 attributes spanning geography, climate (Rainfall, Temp, Humidity, Sunlight), soil health (pH, Moisture, NPK), agricultural inputs (Fertilizer, Pesticide, Seed Score, Irrigation Method), and financial outputs (Yield, Total Production, Market Price, Costs, Revenue, Profit)."),
        ("2. Data Quality Assurance & Precise Imputation:", 
         "Resolved missing values with mathematical accuracy: imputed 32 missing Yield entries using the exact physical relation Yield = Production / Area; imputed Rainfall and Soil Moisture using seasonal-regional grouped medians; verified financial consistency (Revenue = Production \u00d7 Price; Profit = Revenue - Cost)."),
        ("3. Multi-Dimensional Analytics Framework:", 
         "Executed an end-to-end analytical pipeline: (a) Univariate & Bivariate EDA, (b) Climatic regime clustering and pest/disease correlation modeling, (c) Resource & irrigation water efficiency indexing, (d) Econometric analysis & One-Way ANOVA hypothesis testing, and (e) Formulating actionable agronomic recommendations.")
    ]
    for idx, (head, body) in enumerate(sections):
        p = tf3.paragraphs[0] if idx == 0 else tf3.add_paragraph()
        p.text = head
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_RED
        p.space_after = Pt(2)
        
        pb = tf3.add_paragraph()
        pb.text = body
        pb.font.size = Pt(12)
        pb.font.color.rgb = COLOR_DARK
        pb.space_after = Pt(12)
    print("Slide 3 formatted.")

    # -------------------------------------------------------------
    # SLIDE 4: WHO ARE THE END USERS?
    # -------------------------------------------------------------
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.paragraphs[0].text = "WHO ARE THE END USERS?"
            shape.text_frame.paragraphs[0].font.size = Pt(30)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
        elif shape.name == 'Text Placeholder 1' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            shape.left = Inches(0.8)
            shape.top = Inches(1.5)
            shape.width = Inches(11.7)
            shape.height = Inches(4.9)
            
            users = [
                ("1. Farmers & Farmer Producer Organizations (FPOs):",
                 "Make data-backed decisions on optimal crop selection per season, shift from flood to micro-irrigation, and eliminate loss-making summer practices."),
                ("2. Agricultural Extension Officers & Agronomists:",
                 "Issue early pest warning advisories during humid Kharif conditions (54.5% risk) and guide soil moisture conservation practices in dry Zaid months."),
                ("3. Policymakers & State Agricultural Departments:",
                 "Design targeted MSP pricing schemes, direct micro-irrigation subsidies toward water-stressed regions, and plan regional water budgeting."),
                ("4. Agribusinesses, Commodity Traders & Supply Chain Planners:",
                 "Forecast seasonal crop production volumes, schedule procurement timelines, and optimize warehouse cold storage capacity."),
                ("5. Rural Financial Institutions & Crop Insurance Underwriters:",
                 "Utilize empirical climate risk distributions and historical seasonal failure rates for precision actuarial risk and insurance pricing.")
            ]
            for idx, (title, desc) in enumerate(users):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = title
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = COLOR_RED
                p.space_after = Pt(1)
                
                pd = tf.add_paragraph()
                pd.text = "   " + desc
                pd.font.size = Pt(11.5)
                pd.font.color.rgb = COLOR_DARK
                pd.space_after = Pt(8)
    print("Slide 4 formatted.")

    # -------------------------------------------------------------
    # SLIDE 5: Technology Used
    # -------------------------------------------------------------
    slide5 = prs.slides[4]
    for shape in slide5.shapes:
        if shape.name == 'Title 8' and shape.has_text_frame:
            shape.text_frame.paragraphs[0].text = "Technology Stack & Analytical Tools"
            shape.text_frame.paragraphs[0].font.size = Pt(30)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.6)
            shape.width = Inches(10)
        elif shape.name == 'Text Placeholder 6' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            shape.left = Inches(2.2) # Avoid the left template graphic
            shape.top = Inches(1.5)
            shape.width = Inches(10.3)
            shape.height = Inches(5.0)
            
            tech_stack = [
                ("Python 3.11: ", "Core scientific programming runtime for scalable numerical analysis and pipeline automation."),
                ("Pandas & NumPy: ", "Data ingestion, multi-table joins, grouped aggregations, vector math, and mathematical missing value imputation."),
                ("Matplotlib & Seaborn: ", "Multi-panel dashboard generation, bivariate regression plots, KDE density contours, and publication-ready charts."),
                ("SciPy Stats: ", "Statistical inference, Pearson correlation coefficients, and One-Way ANOVA hypothesis testing."),
                ("Jupyter Notebook & VS Code: ", "Interactive exploratory data analysis, code documentation, and inline markdown narration."),
                ("python-pptx & Git/GitHub: ", "Programmatic slide deck synthesis and distributed open-source code version management.")
            ]
            for idx, (name, role) in enumerate(tech_stack):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = BULLET + name
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = COLOR_RED
                p.space_after = Pt(1)
                
                pd = tf.add_paragraph()
                pd.text = "   " + role
                pd.font.size = Pt(11.5)
                pd.font.color.rgb = COLOR_DARK
                pd.space_after = Pt(8)
    print("Slide 5 formatted.")

    # -------------------------------------------------------------
    # SLIDE 6: RESULTS: Executive Summary & Dashboard
    # -------------------------------------------------------------
    slide6 = prs.slides[5]
    for shape in slide6.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.paragraphs[0].text = "RESULTS: Executive Overview & Seasonal KPI Dashboard"
            shape.text_frame.paragraphs[0].font.size = Pt(26)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.5)
            shape.width = Inches(11.5)
        elif 'Text Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            
    if os.path.exists('assets/chart0_overview.png'):
        slide6.shapes.add_picture('assets/chart0_overview.png', Inches(0.8), Inches(1.4), Inches(11.7), Inches(3.2))
        
    tb6 = slide6.shapes.add_textbox(Inches(0.8), Inches(4.75), Inches(11.7), Inches(1.8))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    
    findings = [
        ("Kharif Season (Monsoon): ", "Drives highest productivity (5.64 t/ha avg yield) and revenue (\u20b97.11L), yielding \u20b91.79L net profit. However, humid conditions trigger the highest pest/disease risk (54.5%), necessitating active pest control."),
        ("Rabi Season (Winter): ", "Exhibits stable, reliable performance (5.08 t/ha yield, \u20b987.7k net profit) with balanced input costs (\u20b95.14L) and moderate disease risk (40.5%)."),
        ("Zaid Season (Summer): ", "Experiences intense heat (31.0\u00b0C) and water deficit (299mm rain). Without micro-irrigation, farms suffer negative returns (-\u20b924.8k avg loss) and degraded water efficiency (4.41 t/1000m\u00b3).")
    ]
    for idx, (head, text) in enumerate(findings):
        p = tf6.paragraphs[0] if idx == 0 else tf6.add_paragraph()
        p.text = BULLET + head + text
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)
    print("Slide 6 formatted.")

    # -------------------------------------------------------------
    # HELPER FOR RESULTS SLIDES 7-10 (Code + Chart + Takeaways)
    # -------------------------------------------------------------
    def format_results_slide(slide_idx, title_text, code_path, chart_path, takeaways):
        slide = prs.slides[slide_idx]
        for shape in slide.shapes:
            if shape.name == 'Title 3' and shape.has_text_frame:
                shape.text_frame.paragraphs[0].text = title_text
                shape.text_frame.paragraphs[0].font.size = Pt(24)
                shape.text_frame.paragraphs[0].font.bold = True
                shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
                shape.left = Inches(0.8)
                shape.top = Inches(0.5)
                shape.width = Inches(11.5)
            elif 'Text Placeholder' in shape.name and shape.has_text_frame:
                shape.text_frame.clear()
                
        # Insert Code Image
        if os.path.exists(code_path):
            slide.shapes.add_picture(code_path, Inches(0.8), Inches(1.35), Inches(5.6), Inches(3.45))
            
        # Insert Chart Image
        if os.path.exists(chart_path):
            slide.shapes.add_picture(chart_path, Inches(6.6), Inches(1.35), Inches(5.9), Inches(3.45))
            
        # Insert Bottom Insight Card
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(4.95), Inches(11.7), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for idx, (point_title, point_desc) in enumerate(takeaways):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = BULLET + point_title + ": " + point_desc
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK
            p.space_after = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 7: RESULTS: Seasonal Crop Yield & Production Dynamics
    # -------------------------------------------------------------
    format_results_slide(
        6,
        "RESULTS: Seasonal Crop Yield & Production Dynamics",
        "assets/code_snippet_1.png",
        "assets/chart1_yield_production.png",
        [
            ("Monsoon Volume Dominance", "Kharif season accounts for 48% of total agricultural production volume across all farms, bolstered by widespread monsoon rainfall (852 mm avg)."),
            ("High-Yield Commercial Crops", "Sugarcane achieves standout yields (46.9 t/ha), while cash crops like Chilli, Groundnut, and Pulses show remarkable yield consistency across seasons."),
            ("Staple Crop Variability", "Rice yields peak in Kharif (2.44 t/ha) and drop in summer Zaid, whereas Maize exhibits strong multi-season adaptability across both Kharif and Rabi.")
        ]
    )
    print("Slide 7 formatted.")

    # -------------------------------------------------------------
    # SLIDE 8: RESULTS: Environmental Drivers & Pest/Disease Risk
    # -------------------------------------------------------------
    format_results_slide(
        7,
        "RESULTS: Environmental Drivers & Pest/Disease Risk",
        "assets/code_snippet_2.png",
        "assets/chart2_environmental_disease.png",
        [
            ("Humidity-Pest Outbreak Nexus", "Atmospheric humidity (>70%) exhibits a strong positive correlation with crop disease and pest incidence (r = +0.48, p < 1e-10)."),
            ("Critical Monsoon Risk Window", "Kharif experiences peak disease risk (54.5% avg), requiring early prophylactic bio-pesticide interventions and active field scouting."),
            ("Summer Climate Stress", "Zaid records the lowest humidity and disease incidence (38.2%), yet extreme heat (31.0\u00b0C) elevates evapotranspiration, suppressing crop productivity.")
        ]
    )
    print("Slide 8 formatted.")

    # -------------------------------------------------------------
    # SLIDE 9: RESULTS: Resource Utilization & Irrigation Efficiency
    # -------------------------------------------------------------
    format_results_slide(
        8,
        "RESULTS: Resource Utilization & Irrigation Efficiency",
        "assets/code_snippet_3.png",
        "assets/chart3_irrigation_efficiency.png",
        [
            ("Micro-Irrigation Superiority", "Drip irrigation generates the highest water efficiency (6.27 t/1000m\u00b3) and top profitability (\u20b92.35L avg), delivering high crop output per drop."),
            ("The Flood Irrigation Trap", "Consumes 8,026 m\u00b3 per farm (35% more water than drip) while generating lower yields, causing severe economic losses (-\u20b969.8k avg) in Zaid."),
            ("Rainfed Agriculture Limitations", "Rainfed cultivation thrives during Kharif (+\u20b91.59L profit), but causes catastrophic losses (-\u20b979.5k avg) when attempted during summer Zaid.")
        ]
    )
    print("Slide 9 formatted.")

    # -------------------------------------------------------------
    # SLIDE 10: RESULTS: Economic Outcomes & Profitability Dynamics
    # -------------------------------------------------------------
    format_results_slide(
        9,
        "RESULTS: Economic Outcomes & Profitability Dynamics",
        "assets/code_snippet_4.png",
        "assets/chart4_economic_profitability.png",
        [
            ("Statistical Significance (ANOVA)", "One-Way ANOVA confirms significant net profit differences across seasons (F = 34.29, p = 1.71e-15; reject null hypothesis)."),
            ("Commercial High-Margin Outperformers", "Chilli and Sugarcane produce high net profits across all seasons (>\u20b94.5L to \u20b910.0L/farm), supported by strong market price realizations."),
            ("Staple Crop Cost Squeeze", "Wheat and summer crops face compressed margins due to rigid input and irrigation costs relative to prevailing market prices.")
        ]
    )
    print("Slide 10 formatted.")

    # -------------------------------------------------------------
    # SLIDE 11: Future scope
    # -------------------------------------------------------------
    slide11 = prs.slides[10]
    for shape in slide11.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.paragraphs[0].text = "Future Scope & Strategic Recommendations"
            shape.text_frame.paragraphs[0].font.size = Pt(28)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.5)
            shape.width = Inches(11.5)
        elif 'Text Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            
    tb11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.0))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    
    scopes = [
        ("1. Machine Learning Predictive Modeling: ", 
         "Train Random Forest and XGBoost regressors to predict seasonal crop yields and classify farm profitability based on hyper-local climatic and soil variables."),
        ("2. IoT & Smart Sensor Precision Irrigation: ", 
         "Integrate automated soil moisture probes and weather telemetry to dynamically regulate drip irrigation schedules, conserving 30-40% water."),
        ("3. Climate-Resilient Cropping Architecture: ", 
         "Transition summer (Zaid) cropping toward drought-tolerant pulses, millets, and oilseeds under micro-irrigation, phasing out summer flood farming."),
        ("4. Digital Mandi Linkages & Real-Time Price Discovery: ", 
         "Deploy farmer-facing advisory apps providing real-time mandi prices and Kharif pest advisories to optimize harvesting and sales timing."),
        ("5. Institutional Policy Support & Subsidies: ", 
         "Channel state subsidies into micro-irrigation systems (drip/sprinkler) and establish parametric seasonal crop insurance based on weather indices.")
    ]
    for idx, (title, body) in enumerate(scopes):
        p = tf11.paragraphs[0] if idx == 0 else tf11.add_paragraph()
        p.text = title
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_RED
        p.space_after = Pt(1)
        
        pd = tf11.add_paragraph()
        pd.text = "   " + body
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = COLOR_DARK
        pd.space_after = Pt(7)
    print("Slide 11 formatted.")

    # -------------------------------------------------------------
    # SLIDE 12: GitHub Link
    # -------------------------------------------------------------
    slide12 = prs.slides[11]
    for shape in slide12.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.paragraphs[0].text = "GitHub Repository & Project Deliverables"
            shape.text_frame.paragraphs[0].font.size = Pt(28)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.5)
            shape.width = Inches(11.5)
        elif shape.name == 'TextBox 2' and shape.has_text_frame:
            shape.left = Inches(0.8)
            shape.top = Inches(1.4)
            shape.width = Inches(11.7)
            shape.height = Inches(0.8)
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Project GitHub Repository:"
            p0.font.bold = True
            p0.font.size = Pt(14)
            p0.font.color.rgb = COLOR_DARK
            p1 = tf.add_paragraph()
            p1.text = "https://github.com/manaalsaaas-afk/vois-internship.git"
            p1.font.bold = True
            p1.font.size = Pt(15)
            p1.font.color.rgb = COLOR_RED
        elif 'Text Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            
    tb12 = slide12.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(4.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    
    repo_details = [
        (BULLET + "Repository Architecture: ", "Structured into modular folders: data/, notebooks/, presentation/, and assets/ for full production readiness."),
        (BULLET + "Jupyter Notebook: ", "Seasonal_Agriculture_Performance_Analysis.ipynb with complete end-to-end reproducible analysis, EDA, and statistical tests."),
        (BULLET + "Datasets: ", "Includes both raw (seasonal_agriculture_performance_dataset.csv) and cleaned/imputed datasets with full validation logs."),
        (BULLET + "Presentation Deck: ", "Official VOIS submission deck (VOIS_Major_Project_Seasonal_Agriculture_Performance_Analysis.pptx)."),
        (BULLET + "Visual Artifacts: ", "High-resolution publication-quality figures, charts, and code cards exported at 300 DPI."),
        (BULLET + "Reproducibility: ", "Includes requirements.txt and detailed README.md instructions for one-click environment replication.")
    ]
    for idx, (title, desc) in enumerate(repo_details):
        p = tf12.paragraphs[0] if idx == 0 else tf12.add_paragraph()
        p.text = title + desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(8)
    print("Slide 12 formatted.")

    # -------------------------------------------------------------
    # SLIDE 13: VOIS Course Completion Certificate
    # -------------------------------------------------------------
    slide13 = prs.slides[12]
    for shape in slide13.shapes:
        if shape.name == 'Title 3' and shape.has_text_frame:
            shape.text_frame.clear()
            p0 = shape.text_frame.paragraphs[0]
            p0.text = "VOIS Course Completion Certificate: Data Visualization"
            p0.font.size = Pt(26)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_RED
            shape.left = Inches(0.8)
            shape.top = Inches(0.5)
            shape.width = Inches(11.5)
        elif 'Text Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            
    # Add Certificate Frame Box (Left side)
    rect = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(6.0), Inches(4.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = COLOR_CARD_BG
    rect.line.color.rgb = COLOR_RED
    rect.line.width = Pt(2)
    tf_rect = rect.text_frame
    tf_rect.word_wrap = True
    p_box = tf_rect.paragraphs[0]
    p_box.text = "\n\n\n\n[ Official VOIS Course Completion Certificate ]\n\nCourse Name: Data Visualization\nIssued by: Vodafone Intelligent Solutions (VOIS) & AICTE\n\n(Attach your certificate image / PDF scan here)"
    p_box.alignment = PP_ALIGN.CENTER
    p_box.font.size = Pt(13)
    p_box.font.color.rgb = COLOR_MUTED
    
    # Add Competencies Box (Right side)
    tb13 = slide13.shapes.add_textbox(Inches(7.1), Inches(1.5), Inches(5.4), Inches(4.7))
    tf13 = tb13.text_frame
    tf13.word_wrap = True
    
    comps = [
        ("Program Track: ", "VOIS AICTE Emerging Technologies Internship (Batch 1, 2026-2027)"),
        ("Course Completed: ", "Data Visualization with Python"),
        ("Key Competencies Acquired: ", ""),
        (BULLET + "Exploratory Visual Analytics: ", "Multi-dimensional feature mapping, KDE distributions, and correlation matrices."),
        (BULLET + "Scientific Plotting: ", "Mastery of Matplotlib and Seaborn for publication-grade charts and statistical annotations."),
        (BULLET + "Executive Data Storytelling: ", "Translating complex agricultural metrics into intuitive decision-making dashboards."),
        (BULLET + "Business Intelligence: ", "Designing KPI scorecards and stakeholder-oriented analytical reports.")
    ]
    for idx, (title, desc) in enumerate(comps):
        p = tf13.paragraphs[0] if idx == 0 else tf13.add_paragraph()
        p.text = title + desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_RED if "Track" in title or "Course" in title or "Acquired" in title else COLOR_DARK
        p.font.bold = True if "Track" in title or "Course" in title or "Acquired" in title else False
        p.space_after = Pt(6)
    print("Slide 13 formatted.")

    # -------------------------------------------------------------
    # SLIDE 14: Thank You
    # -------------------------------------------------------------
    slide14 = prs.slides[13]
    for shape in slide14.shapes:
        if shape.name == 'Title 1' and shape.has_text_frame:
            shape.text_frame.clear()
            p0 = shape.text_frame.paragraphs[0]
            p0.text = "Thank You!"
            p0.font.size = Pt(40)
            p0.font.bold = True
            p0.font.color.rgb = COLOR_RED
            p0.alignment = PP_ALIGN.CENTER
            shape.left = Inches(1.5)
            shape.top = Inches(1.8)
            shape.width = Inches(10.3)
            shape.height = Inches(1.0)
        elif 'Text Placeholder' in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            
    tb14 = slide14.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(10.3), Inches(3.0))
    tf14 = tb14.text_frame
    tf14.word_wrap = True
    
    p1 = tf14.paragraphs[0]
    p1.text = "Seasonal Agriculture Performance Analysis"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_DARK
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(6)
    
    p2 = tf14.add_paragraph()
    p2.text = "VOIS AICTE Internship Major Project (Batch 1, 2026-2027)"
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_MUTED
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(20)
    
    p3 = tf14.add_paragraph()
    p3.text = "Special thanks to Vodafone Intelligent Solutions (VOIS) and the AICTE Internship Team\nfor their invaluable guidance, curriculum, and mentorship."
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_DARK
    p3.alignment = PP_ALIGN.CENTER
    p3.space_after = Pt(20)
    
    p4 = tf14.add_paragraph()
    p4.text = "Questions & Discussions are Welcome"
    p4.font.size = Pt(16)
    p4.font.bold = True
    p4.font.color.rgb = COLOR_RED
    p4.alignment = PP_ALIGN.CENTER
    print("Slide 14 formatted.")

    prs.save(output_path)
    print(f"Successfully generated presentation: {output_path}")

if __name__ == '__main__':
    build_presentation()
